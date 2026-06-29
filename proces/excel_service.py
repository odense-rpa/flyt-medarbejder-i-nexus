import openpyxl
from pathlib import Path
from proces.models import Medarbejder, Opgavetype
import logging


def afkode_excel(excel_sti: str | Path) -> list[Medarbejder]:
    wb = openpyxl.load_workbook(excel_sti, data_only=True)
    medarbejdere: list[Medarbejder] = []

    for opgavetype in Opgavetype:
        if opgavetype.value not in wb.sheetnames:
            continue

        ws = wb[opgavetype.value]
        headers = [cell.value for cell in ws[1]]

        medarbejder_col = headers.index("Medarbejdere")

        if opgavetype in (Opgavetype.TILFOEJELSE, Opgavetype.FJERNELSE):
            kontekst_col = headers.index("Organisation")
            kontekst_felt = "organisation"
        else:
            kontekst_col = headers.index("Rolle")
            kontekst_felt = "rolle"

        current_kontekst: str | None = None

        for row in ws.iter_rows(min_row=2, values_only=True):
            kontekst_value = row[kontekst_col]
            initialer_value = row[medarbejder_col]

            if kontekst_value:
                current_kontekst = str(kontekst_value).strip()

            if not initialer_value:
                continue

            for initialer in str(initialer_value).splitlines():
                initialer = initialer.strip()
                if not initialer:
                    continue

                m = Medarbejder(
                    initialer=initialer,
                    opgavetype=opgavetype,
                    **{kontekst_felt: current_kontekst},
                )
                medarbejdere.append(m)

    return medarbejdere

def find_excel_fil(mappe_sti) -> str | Path:
    logger = logging.getLogger(__name__)
    
    over_mappe = Path(mappe_sti)
    if not over_mappe.exists() or not over_mappe.is_dir():
        raise ValueError(f"Mappen findes ikke eller er ikke en mappe: {over_mappe}")
    
    # Find only direct child folders inside the provided path (not the path itself).
    mapper = [sti for sti in over_mappe.iterdir() if sti.is_dir()]
    
    if len(mapper) > 1:
        logger.error("Mere end en mappe fundet, smider fejl og stopper")
        raise ValueError("Mere end en opgavemappe i hovedmappe")
    
    if len(mapper) == 0:
        logger.info("Ingen mappe fundet, går ud fra der ikke ligger nogle opgaver, stopper")
        return ""
    
    logger.info("Kun en mappe fundet")
    mappe = mapper[0]
    
    teskt_fil = [*mappe.glob("*.txt")]
    
    if len(teskt_fil) >= 1:
        logger.info("Tekst fil fundet, derfor går vi udfra at opgaveløsning er igang/afsluttet på dette opgavesæt")
        return ""
    
    excel_fil = [*mappe.glob("*.xlsx"), *mappe.glob("*.xlsm")]
    if not excel_fil:
        return ""
    if len(excel_fil) > 1:
        raise ValueError("Mere end 1 excel-fil fundet")
    return excel_fil[0]

def find_tekst_fil(mappe_sti) -> Path | None:
    over_mappe = Path(mappe_sti)
    if not over_mappe.exists() or not over_mappe.is_dir():
        raise ValueError(f"Mappen findes ikke eller er ikke en mappe: {over_mappe}")

    
    # Find only direct child folders inside the provided path (not the path itself).
    mapper = [sti for sti in over_mappe.iterdir() if sti.is_dir()]
    
    if len(mapper) > 1:
        raise ValueError("Mere end en opgavemappe i hovedmappe")
    
    mappe = mapper[0]
    
    teskt_fil = [*mappe.glob("*.txt")]
    
    if len(teskt_fil) >= 1:
        return teskt_fil[0]
    else:
        return None